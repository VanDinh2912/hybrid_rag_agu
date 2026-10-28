"""Script xuất tài liệu kiến trúc hệ thống sang định dạng TXT và DOCX.

Sử dụng:
    python scripts/export_architecture.py
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

# Đọc file markdown
md_path = Path("docs", "architecture", "SYSTEM_ARCHITECTURE.md")
md_content = md_path.read_text(encoding="utf-8")

# Xuất sang TXT
txt_path = Path("docs", "architecture", "SYSTEM_ARCHITECTURE.txt")
txt_path.write_text(md_content, encoding="utf-8")

# Xuất sang DOCX
doc = Document()

# Thiết lập font mặc định
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(11)

# Tiêu đề chính
title = doc.add_heading('Kiến trúc Hệ thống Hybrid RAG - Tra cứu Quy chế Đào tạo ĐHAG', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Parse markdown và thêm vào DOCX
lines = md_content.split('\n')
i = 0
while i < len(lines):
    line = lines[i]
    
    if line.startswith('# '):
        # H1
        doc.add_heading(line[2:], 0)
    elif line.startswith('## '):
        # H2
        doc.add_heading(line[3:], 1)
    elif line.startswith('### '):
        # H3
        doc.add_heading(line[4:], 2)
    elif line.startswith('#### '):
        # H4
        doc.add_heading(line[5:], 3)
    elif line.startswith('```'):
        # Code block - thêm như paragraph với font monospace
        i += 1
        code_lines = []
        while i < len(lines) and not lines[i].startswith('```'):
            code_lines.append(lines[i])
            i += 1
        p = doc.add_paragraph('\n'.join(code_lines))
        p.style = 'No Spacing'
        for run in p.runs:
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(50, 50, 50)
    elif line.startswith('- '):
        # Bullet list
        p = doc.add_paragraph(line[2:], style='List Bullet')
    elif line.startswith('|') and line.endswith('|'):
        # Table markdown - simplified: chỉ xử lý dòng hiện tại
        row = [cell.strip() for cell in line.split('|')[1:-1]]
        # Nếu row có chứa "---" thì bỏ qua (separator)
        if not all(c.startswith('-') for c in row):
            p = doc.add_paragraph(' | '.join(row))
            p.runs[0].font.size = Pt(9)
    elif line.strip() == '':
        # Empty line
        doc.add_paragraph()
    else:
        # Paragraph thường
        p = doc.add_paragraph(line)
        for run in p.runs:
            run.font.size = Pt(11)
    
    i += 1

# Lưu DOCX
docx_path = Path("docs", "architecture", "SYSTEM_ARCHITECTURE.docx")
doc.save(docx_path)

# Tạo file báo cáo (nơi duy nhất được phép ghi log: logs/export/)
report_path = Path("logs", "export", "export_report.txt")
report_path.parent.mkdir(parents=True, exist_ok=True)
report = f"""Export Report
==============

Source: docs/architecture/SYSTEM_ARCHITECTURE.md
Output files:
  - TXT: {txt_path.absolute()}
  - DOCX: {docx_path.absolute()}

Status: Success
"""
report_path.write_text(report, encoding="utf-8")
