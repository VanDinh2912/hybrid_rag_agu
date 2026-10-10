"""Đọc tài liệu quy chế đa định dạng (.txt, .pdf, .docx) đệ quy từ thư mục corpus.

PDF scan ảnh (pypdf trích ra 0 ký tự) được tự động OCR bằng pipeline thuần pip:
pypdfium2 render từng trang thành ảnh -> rapidocr-onnxruntime nhận diện tiếng Việt.
Không phải cài phần mềm hệ thống (Tesseract/Poppler).

Kết quả OCR được cache theo (đường dẫn, kích thước, mtime, dpi, ngưỡng) ra
./.ocr_cache/ vì một lần OCR toàn bộ Tầng 1 mất ~10 phút, trong khi learn.py
rebuild lại collection sau mỗi bài học.
"""

import hashlib
import json
import os
import re
import time

import numpy as np
import pypdfium2 as pdfium
from docx import Document as DocxDocument
from pypdf import PdfReader

# Trang ít ký tự hơn ngưỡng này bị coi là trang scan và đưa sang OCR: văn bản hành
# chính một trang thường > 1.000 ký tự, trang scan trả về 0, nên ngưỡng thấp vẫn an
# toàn mà bắt được cả trang chỉ có con dấu/số trang.
DEFAULT_MIN_CHARS_PER_PAGE = 50
# 200 dpi đo trên corpus thật cho chất lượng ngang 300 dpi (model OCR là nút thắt,
# không phải độ phân giải) nhưng nhanh hơn ~15%.
DEFAULT_OCR_DPI = 200
DEFAULT_CACHE_DIR = "./.ocr_cache"
OCR_INSTALL_HINT = "pip install pypdfium2 rapidocr-onnxruntime Pillow"

# Nạp model ONNX một lần cho cả tiến trình: dựng engine mất ~1s và 13 file scan
# dùng chung một bộ model.
_OCR_ENGINE = None


def _extract_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


def _extract_docx(path: str) -> str:
    doc = DocxDocument(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


TEXT_EXTRACTORS = {".txt": _extract_txt, ".docx": _extract_docx}

# Định dạng nhận diện được nhưng không trích xuất nổi — phải báo rõ thay vì bỏ qua
# âm thầm, vì đây là dữ liệu đề tài cần dùng.
UNSUPPORTED_FORMATS = {
    ".doc": "Word 97-2003 (.doc) — python-docx chỉ đọc .docx, hãy Save As sang .docx",
}


def _get_ocr_engine():
    """Khởi tạo (một lần) engine RapidOCR; báo lỗi kèm lệnh cài nếu thiếu thư viện."""
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
        except ImportError as exc:
            raise RuntimeError(f"thiếu thư viện OCR ({exc}) — cài bằng: {OCR_INSTALL_HINT}") from exc
        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def _ocr_page_text(engine, image) -> str:
    """OCR một ảnh trang, trả về các dòng text theo thứ tự đọc (mỗi dòng một '\\n')."""
    result, _elapse = engine(np.array(image))
    if not result:
        return ""
    # result: [[box, text, score], ...] — giữ thứ tự box mà detector trả về (đã theo
    # thứ tự đọc từ trên xuống), nối bằng '\n' để mốc "Dieu N." vẫn nằm ở đầu dòng.
    return "\n".join(item[1].strip() for item in result if item[1].strip())


class DocumentLoader:
    """Đọc tài liệu quy chế đa định dạng đệ quy, kèm OCR cho PDF scan.

    Mỗi document kèm metadata phân tầng: tier = tên thư mục con cấp 1 (tính từ
    data_dir) chứa file đó, ví dụ "Tầng 1", "Tầng 2"; file nằm ngay trong data_dir
    nhận default_tier.

    File không đọc được (.doc, file hỏng, scan mà OCR thất bại) được tích lũy vào
    self.skipped dạng [(đường_dẫn, lý_do)]; thống kê OCR vào self.ocr_stats dạng
    [(đường_dẫn, số_trang_ocr, tổng_số_trang, số_ký_tự)] — để tiến trình nạp in báo
    cáo tổng kết thay vì thất thoát dữ liệu trong im lặng.
    """

    def __init__(self, data_dir: str = "./data/corpus", default_tier: str = "Khác",
                 enable_ocr: bool | None = None, ocr_dpi: int | None = None,
                 min_chars_per_page: int | None = None, ocr_cache_dir: str | None = None):
        self.data_dir = data_dir
        self.default_tier = default_tier
        self.enable_ocr = (
            os.getenv("OCR_ENABLED", "1").strip().lower() not in {"0", "false", "no", "off"}
            if enable_ocr is None else enable_ocr
        )
        self.ocr_dpi = int(os.getenv("OCR_DPI", DEFAULT_OCR_DPI)) if ocr_dpi is None else ocr_dpi
        self.min_chars_per_page = (
            int(os.getenv("OCR_MIN_CHARS_PER_PAGE", DEFAULT_MIN_CHARS_PER_PAGE))
            if min_chars_per_page is None else min_chars_per_page
        )
        self.ocr_cache_dir = os.getenv("OCR_CACHE_DIR", DEFAULT_CACHE_DIR) if ocr_cache_dir is None else ocr_cache_dir
        self.skipped: list[tuple[str, str]] = []
        self.ocr_stats: list[tuple[str, int, int, int]] = []
        os.makedirs(self.data_dir, exist_ok=True)

    def _tier_of(self, path: str) -> str:
        rel_parts = os.path.relpath(path, self.data_dir).split(os.sep)
        return rel_parts[0] if len(rel_parts) > 1 else self.default_tier

    def _skip(self, path: str, reason: str) -> None:
        self.skipped.append((path, reason))
        print(f"[Cảnh báo] Bỏ qua {path}: {reason}")

    # ---------------------------------------------------------------- OCR cache

    def _cache_file(self, path: str) -> str:
        """Tên file cache gắn với nội dung + tham số OCR, để sửa file là cache tự hỏng."""
        stat = os.stat(path)
        key = (f"{os.path.relpath(path, self.data_dir)}|{stat.st_size}|{int(stat.st_mtime)}"
               f"|dpi={self.ocr_dpi}|min={self.min_chars_per_page}")
        digest = hashlib.sha1(key.encode("utf-8")).hexdigest()[:16]
        stem = re.sub(r"[^\w.-]+", "_", os.path.splitext(os.path.basename(path))[0])[:60]
        return os.path.join(self.ocr_cache_dir, f"{stem}__{digest}.json")

    def _cache_load(self, path: str, scan_pages: list[int]) -> dict[int, str] | None:
        cache_file = self._cache_file(path)
        if not os.path.exists(cache_file):
            return None
        try:
            with open(cache_file, "r", encoding="utf-8") as f:
                pages = {int(k): v for k, v in json.load(f).get("pages", {}).items()}
        except (OSError, ValueError):
            return None
        return pages if set(pages) == set(scan_pages) else None

    def _cache_save(self, path: str, pages: dict[int, str]) -> None:
        os.makedirs(self.ocr_cache_dir, exist_ok=True)
        with open(self._cache_file(path), "w", encoding="utf-8") as f:
            json.dump({"pages": {str(k): v for k, v in pages.items()}}, f, ensure_ascii=False)

    # ---------------------------------------------------------------- Trích xuất

    def _ocr_pages(self, path: str, scan_pages: list[int], total_pages: int) -> dict[int, str]:
        """OCR các trang scan của một PDF, trả về {số_thứ_tự_trang: text}."""
        rel = os.path.relpath(path, self.data_dir)
        cached = self._cache_load(path, scan_pages)
        if cached is not None:
            chars = sum(len(text) for text in cached.values())
            self.ocr_stats.append((rel, len(cached), total_pages, chars))
            print(f"[OCR] Dùng lại cache: {rel} ({len(cached)} trang, {chars:,} ký tự)")
            return cached

        print(f"[OCR] Đang xử lý file scan: {rel} "
              f"({len(scan_pages)}/{total_pages} trang cần OCR, dpi={self.ocr_dpi})")
        engine = _get_ocr_engine()
        started = time.perf_counter()
        pages: dict[int, str] = {}
        with pdfium.PdfDocument(path) as pdf:
            for index in scan_pages:
                page_started = time.perf_counter()
                image = pdf[index].render(scale=self.ocr_dpi / 72).to_pil().convert("RGB")
                text = _ocr_page_text(engine, image)
                pages[index] = text
                print(f"[OCR]   trang {index + 1}/{total_pages}: "
                      f"{len(text):,} ký tự ({time.perf_counter() - page_started:.1f}s)")

        chars = sum(len(text) for text in pages.values())
        self._cache_save(path, pages)
        self.ocr_stats.append((rel, len(pages), total_pages, chars))
        print(f"[OCR] Hoàn tất trích xuất thành công: {rel} — "
              f"{len(pages)} trang, {chars:,} ký tự trong {time.perf_counter() - started:.0f}s (đã cache)")
        return pages

    def _extract_pdf(self, path: str) -> str | None:
        """Trích xuất PDF theo từng trang; trang nghèo chữ (scan ảnh) thì OCR.

        Xét theo trang chứ không theo file: một PDF có thể vừa có trang text vừa có
        trang chụp ảnh (bìa, phụ lục), OCR cả file sẽ tốn thời gian vô ích.
        Trả về None khi đã tự ghi nhận vào self.skipped.
        """
        reader = PdfReader(path)
        pages = [page.extract_text() or "" for page in reader.pages]
        scan_pages = [i for i, text in enumerate(pages)
                      if len(text.strip()) < self.min_chars_per_page]
        if not scan_pages:
            return "\n".join(pages)

        if not self.enable_ocr:
            self._skip(path, f"{len(scan_pages)}/{len(pages)} trang là scan ảnh "
                             "nhưng OCR đang tắt (đặt OCR_ENABLED=1 để bật)")
            return None

        ocr_pages = self._ocr_pages(path, scan_pages, len(pages))
        for index, text in ocr_pages.items():
            pages[index] = text
        return "\n".join(pages)

    def _read_document(self, path: str) -> str | None:
        ext = os.path.splitext(path)[1].lower()
        try:
            content = self._extract_pdf(path) if ext == ".pdf" else TEXT_EXTRACTORS[ext](path)
        except Exception as exc:
            self._skip(path, f"lỗi khi trích xuất ({exc})")
            return None
        if content is None:
            return None
        if not content.strip():
            self._skip(path, "không trích xuất được chữ (kể cả sau OCR) — file scan mờ/trắng")
            return None
        return content

    def load_documents(self) -> list[dict]:
        """Quét đệ quy data_dir, trả về [{title, content, tier}] của các file hỗ trợ."""
        self.skipped = []
        self.ocr_stats = []
        paths = []
        for root, _dirs, files in os.walk(self.data_dir):
            for filename in files:
                path = os.path.join(root, filename)
                ext = os.path.splitext(filename)[1].lower()
                if ext == ".pdf" or ext in TEXT_EXTRACTORS:
                    paths.append(path)
                elif ext in UNSUPPORTED_FORMATS:
                    self._skip(path, UNSUPPORTED_FORMATS[ext])

        documents = []
        # Sắp xếp để thứ tự tài liệu (và chunk_id sau này) ổn định giữa các lần chạy
        for path in sorted(paths):
            content = self._read_document(path)
            if content is None:
                continue
            documents.append({
                "title": os.path.basename(path),
                "content": content,
                "tier": self._tier_of(path),
            })
        return documents
